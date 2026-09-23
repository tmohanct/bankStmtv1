"""Assemble and style final workbooks from transaction analysis tables."""

import re
from copy import copy
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.export.excel_safety import sanitize_excel_frame, sanitize_excel_value, force_leading_equals_to_text
from src.utils.file_utils import append_timestamp_if_exists
from src.transform.rule_schedule import build_rule_schedule, order_rule_transactions
from src.export.rule_schedule import add_rule_schedule
from src.export.excel_safety import suppress_number_as_text_warnings

# Re-export helpers for callers of the previous builder API.
from src.transform.analysis import (
    _coerce_excel_date,
    _ensure_columns,
    _first_present_column,
    _parse_rule_amount,
    _normalize_rule_text,
    _load_rules,
    _build_text_rule_sheet,
    _build_amount_rule_sheet,
    _merge_rule_sheet_frames,
    _build_rule_sheets,
    _to_numeric,
    _build_cheque_sheet,
    _is_return_reject_detail,
    _build_return_reject_sheet,
    _cheque_sort_value,
    _build_repeat_sheet,
    _build_top_sheet,
    _build_month_dr_cr_sheet,
)
from src.export.monthly_chart import (
    _format_month_dr_cr_chart_label,
    _load_chart_font,
    _draw_dotted_horizontal_line,
    _format_month_dr_cr_axis_label,
    _nice_axis_step,
    _draw_rotated_text,
    _add_month_dr_cr_chart_image,
    MONTH_DR_CR_CHART_IMAGE_SIZE,
    MONTH_DR_CR_DATA_LABEL_FONT_SIZE,
)

FONT_NORMAL = Font(name="Aptos", size=10)
FONT_HEADER = Font(name="Aptos", size=10, bold=True)
FONT_FOOTNOTE = Font(name="Aptos", size=10, italic=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)
HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9E1F2")
ALT_ROW_FILLS = [
    PatternFill(fill_type="solid", fgColor="EAF3FB"),
    PatternFill(fill_type="solid", fgColor="F8F1E5"),
]
MONTH_LABEL_FILL = PatternFill(fill_type="solid", fgColor="C9D5EA")
MONTH_TITLE_FILL = PatternFill(fill_type="solid", fgColor="DDEBF7")
MONTH_VALUE_ROW_FILLS = [
    PatternFill(fill_type="solid", fgColor="F2F2F2"),
    PatternFill(fill_type="solid", fgColor="FFFFFF"),
]
REPEAT_GROUP_FILLS = [
    PatternFill(fill_type="solid", fgColor="FCE4D6"),
    PatternFill(fill_type="solid", fgColor="D9EAD3"),
    PatternFill(fill_type="solid", fgColor="D9E1F2"),
]
PDF_STATUS_SHEET_NAME = "PDF_Status"
PDF_STATUS_COLUMNS = ["PDF", "Check", "Status", "Result", "Details"]
PDF_ACCOUNT_SUMMARY_LABELS = [
    "Customer Name",
    "Bank Name",
    "Account Number",
    "Address",
    "Statement Date Between",
]
PDF_STATUS_TABLE_START_ROW = len(PDF_ACCOUNT_SUMMARY_LABELS) + 3
PDF_STATUS_FILLS = {
    "PASS": PatternFill(fill_type="solid", fgColor="C6EFCE"),
    "WARNING": PatternFill(fill_type="solid", fgColor="FFF2CC"),
    "FAIL": PatternFill(fill_type="solid", fgColor="F4CCCC"),
    "UNASSESSABLE": PatternFill(fill_type="solid", fgColor="D9E1F2"),
}
PDF_STATUS_RANK = {"PASS": 0, "WARNING": 1, "UNASSESSABLE": 2, "FAIL": 3}
INDIAN_NUMBER_FORMAT = "#,##,##0.00"
INDIAN_NUMBER_FORMAT_NO_DECIMAL = "#,##,##0"
DATE_NUMBER_FORMAT = "yyyy-mm-dd"
AMOUNT_COLUMN_WIDTH = 16
AMOUNT_COLUMN_PADDING = 4
PDF_STATUS_MAX_COLUMN_WIDTH = 108
THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
MONTH_DR_CR_FOOTNOTE = "#.OF Dr/Cr & Avg takes only amount Greater than 30. Less than 30 not counted."
FINAL_EXCLUDED_COLUMNS = ("Detail_Clean",)


def _sanitize_excel_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    return sanitize_excel_value(value)


def _sanitize_excel_frame(frame: pd.DataFrame) -> pd.DataFrame:
    return sanitize_excel_frame(frame)


def _sanitize_sheet_name(name: str) -> str:
    safe = re.sub(r"[\\/*?:\[\]]", "_", _sanitize_excel_value(str(name)).strip())
    safe = safe or "Sheet"
    return safe[:31]


def _unique_sheet_name(name: str, used_names: set[str]) -> str:
    base = _sanitize_sheet_name(name)
    candidate = base
    index = 1
    while candidate in used_names:
        suffix = f"_{index}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        index += 1
    used_names.add(candidate)
    return candidate


def _normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _round_money_for_excel(value: Any) -> int | None:
    if value in (None, ""):
        return None

    try:
        if pd.isna(value):
            return None
    except TypeError:
        pass

    try:
        rounded = Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None

    return int(rounded)


def _displayed_amount_length(value: Any) -> int:
    """Length of an integer with Indian digit grouping and an optional minus sign."""
    if isinstance(value, float) and pd.isna(value):
        return 0
    if not isinstance(value, (int, float)):
        return len(str(value or ""))
    digits = str(abs(int(value)))
    separators = 0 if len(digits) <= 3 else 1 + (len(digits) - 4) // 2
    return len(digits) + separators + (value < 0)


def _exclude_final_columns(frame: pd.DataFrame) -> pd.DataFrame:
    """Keep internal matching columns out of every final workbook sheet."""
    return frame.drop(columns=list(FINAL_EXCLUDED_COLUMNS), errors="ignore")


def _exclude_source_column(frame: pd.DataFrame, include_source: bool) -> pd.DataFrame:
    """Hide source bookkeeping from all final sheets for a single-PDF run."""
    if include_source:
        return frame
    return frame.drop(columns=["Source"], errors="ignore")


def _build_pdf_account_summary_rows(
    source_pdf_paths: list[Path] | None,
    source_pdf_passwords: list[str | None] | None = None,
) -> list[tuple[str, str]]:
    from src.utils.pdf_status_reader import read_first_page

    if not source_pdf_paths:
        return [(label, "Source PDF not provided") for label in PDF_ACCOUNT_SUMMARY_LABELS]
    password = source_pdf_passwords[0] if source_pdf_passwords else None
    header = read_first_page(Path(source_pdf_paths[0]), password)
    return [(label, header.values.get(label) or "Not found on first page") for label in PDF_ACCOUNT_SUMMARY_LABELS]


def _build_single_pdf_status_rows(
    pdf_path: Path, password: str | None = None, *, audit: dict[str, object] | None = None
) -> list[dict[str, str]]:
    from src.transform.pdf_status import inspect_pdf

    return inspect_pdf(Path(pdf_path), password, audit=audit)


def _build_pdf_status_sheet(
    source_pdf_paths: list[Path] | None,
    source_pdf_passwords: list[str | None] | None = None,
    source_audits: list[dict[str, object]] | None = None,
) -> pd.DataFrame:
    from src.transform.pdf_status import row

    if not source_pdf_paths:
        return pd.DataFrame([row("", "Overall PDF modification status", "WARNING",
                                 "Source PDF path was not provided",
                                 "Run through the CLI to include source PDF integrity checks.")],
                            columns=PDF_STATUS_COLUMNS)
    rows = []
    for index, path in enumerate(source_pdf_paths):
        password = source_pdf_passwords[index] if source_pdf_passwords and index < len(source_pdf_passwords) else None
        audit = source_audits[index] if source_audits and index < len(source_audits) else {}
        rows.extend(_build_single_pdf_status_rows(Path(path), password, audit=audit))
    if len(source_pdf_paths) > 1:
        rows.insert(0, row("", "Account summary scope", "PASS", "Top summary refers to the first PDF",
                           "Each PDF's first-page fields and modification findings are listed separately below."))
    return pd.DataFrame(rows, columns=PDF_STATUS_COLUMNS)


def _apply_base_style(workbook) -> None:
    left_headers = {"date", "detail", "details", "detailclean", "cheque", "chequeno", "source", "pdf", "check", "result"}
    center_headers = {"sno", "status"}
    right_headers = {"debit", "credit", "balance"}
    date_headers = {"date", "txndate", "valuedate"}
    text_headers = {"cheque", "chequeno"}

    for ws in workbook.worksheets:
        if ws.title.lower() in {"month_dr_cr", PDF_STATUS_SHEET_NAME.lower()}:
            continue

        max_row = ws.max_row
        max_col = ws.max_column

        if max_row < 1 or max_col < 1:
            continue

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        header_to_col: dict[str, int] = {}
        for col_idx in range(1, max_col + 1):
            header_value = ws.cell(row=1, column=col_idx).value
            header_to_col[_normalize_header(header_value)] = col_idx

        align_by_col: dict[int, Alignment] = {}
        for normalized_header, col_idx in header_to_col.items():
            if normalized_header in left_headers:
                align_by_col[col_idx] = ALIGN_LEFT
            elif normalized_header in right_headers:
                align_by_col[col_idx] = ALIGN_RIGHT
            elif normalized_header in center_headers:
                align_by_col[col_idx] = ALIGN_CENTER
            else:
                align_by_col[col_idx] = ALIGN_CENTER

        numeric_cols = {
            col_idx
            for normalized_header, col_idx in header_to_col.items()
            if normalized_header in right_headers
        }
        date_cols = {
            col_idx
            for normalized_header, col_idx in header_to_col.items()
            if normalized_header in date_headers
        }
        text_cols = {
            col_idx
            for normalized_header, col_idx in header_to_col.items()
            if normalized_header in text_headers
        }

        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = FONT_HEADER
            cell.alignment = ALIGN_CENTER
            cell.fill = HEADER_FILL

        for row_idx in range(2, max_row + 1):
            fill = ALT_ROW_FILLS[(row_idx - 2) % 2]
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = FONT_NORMAL
                cell.alignment = align_by_col.get(col_idx, ALIGN_CENTER)
                if col_idx in numeric_cols:
                    cell.number_format = INDIAN_NUMBER_FORMAT_NO_DECIMAL
                if col_idx in date_cols:
                    parsed_date = _coerce_excel_date(cell.value)
                    if parsed_date is not None:
                        cell.value = parsed_date
                    cell.number_format = DATE_NUMBER_FORMAT
                if col_idx in text_cols and cell.value not in (None, ""):
                    cell.value = str(cell.value)
                    cell.number_format = "@"
                cell.fill = fill

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col[: min(len(col), 400)]:
                value = "" if cell.value is None else str(cell.value)
                if len(value) > max_len:
                    max_len = len(value)
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 48)

        for col_idx in numeric_cols:
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[col_letter].width = AMOUNT_COLUMN_WIDTH


def _apply_final_layout(workbook) -> None:
    """Apply the workbook-wide display settings requested for final output."""
    date_headers = {"date", "txndate", "valuedate"}

    for ws in workbook.worksheets:
        if ws.title.lower() == PDF_STATUS_SHEET_NAME.lower():
            continue

        date_columns = {
            col_idx
            for col_idx in range(1, ws.max_column + 1)
            if _normalize_header(ws.cell(row=1, column=col_idx).value) in date_headers
        }
        amount_columns = {
            col_idx
            for col_idx in range(1, ws.max_column + 1)
            if _normalize_header(ws.cell(row=1, column=col_idx).value)
            in {"debit", "credit", "balance"}
        }

        for row_idx in range(1, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = (
                36 if ws.title.lower() == "month_dr_cr" and row_idx == 1 else 19
            )
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                alignment = copy(cell.alignment)
                alignment.wrap_text = False
                cell.alignment = alignment

        if ws.title.lower() == "month_dr_cr":
            continue

        for col_idx in range(1, ws.max_column + 1):
            column_letter = get_column_letter(col_idx)
            if col_idx in date_columns:
                ws.column_dimensions[column_letter].width = 13
                continue

            max_length = 0
            for row_idx in range(1, ws.max_row + 1):
                value = ws.cell(row=row_idx, column=col_idx).value
                if value is None:
                    continue
                display_length = (
                    _displayed_amount_length(value)
                    if col_idx in amount_columns and row_idx > 1
                    else max((len(line) for line in str(value).splitlines()), default=0)
                )
                max_length = max(max_length, display_length)
            if col_idx in amount_columns:
                ws.column_dimensions[column_letter].width = min(
                    max(AMOUNT_COLUMN_WIDTH, max_length + AMOUNT_COLUMN_PADDING), 255
                )
            else:
                # Excel supports widths only up to 255 characters.
                ws.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 255)


def _apply_month_dr_cr_style(workbook, sheet_name: str, customer_name: str) -> None:
    if sheet_name not in workbook.sheetnames:
        return

    ws = workbook[sheet_name]
    if ws.max_row < 1 or ws.max_column < 1:
        return

    data_max_col = ws.max_column
    ws.insert_rows(1)
    data_max_row = ws.max_row

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=data_max_col)
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = _sanitize_excel_value(customer_name)
    title_cell.font = Font(name="Aptos", size=16)
    title_cell.fill = MONTH_TITLE_FILL
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = ws.cell(row=2, column=1).coordinate + ":" + ws.cell(row=data_max_row, column=data_max_col).coordinate

    for row_idx in range(2, data_max_row + 1):
        row_label = str(ws.cell(row=row_idx, column=1).value or "").strip()
        is_total_row = row_label.lower() == "total"
        for col_idx in range(1, data_max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = FONT_NORMAL
            cell.border = THIN_BORDER

            if row_idx == 2:
                cell.font = FONT_HEADER
                cell.alignment = ALIGN_LEFT if col_idx == 1 else ALIGN_CENTER
            elif col_idx == 1:
                cell.font = FONT_HEADER
                cell.alignment = ALIGN_LEFT
                cell.fill = MONTH_LABEL_FILL
            else:
                header_value = str(ws.cell(row=2, column=col_idx).value or "").strip()
                if is_total_row:
                    cell.font = FONT_HEADER
                cell.alignment = ALIGN_RIGHT
                if isinstance(cell.value, (int, float)):
                    cell.number_format = INDIAN_NUMBER_FORMAT_NO_DECIMAL
                cell.fill = MONTH_VALUE_ROW_FILLS[(row_idx - 3) % len(MONTH_VALUE_ROW_FILLS)]

    ws.column_dimensions["A"].width = 12
    width_map = {
        "Dr": 16,
        "Cr": 16,
        "Net": 14,
        "EOM Balance": 16,
        "#.Of.Dr": 10,
        "#.Of.Cr": 10,
        "Avg.Dr": 14,
        "Avg.Cr": 14,
    }
    for col_idx in range(2, data_max_col + 1):
        header_value = str(ws.cell(row=2, column=col_idx).value or "").strip()
        column_letter = ws.cell(row=2, column=col_idx).column_letter
        max_length = max(
            (_displayed_amount_length(ws.cell(row=row_idx, column=col_idx).value)
             for row_idx in range(3, data_max_row + 1)
             if ws.cell(row=row_idx, column=col_idx).value is not None),
            default=len(header_value),
        )
        ws.column_dimensions[column_letter].width = min(
            max(width_map.get(header_value, 14), max_length + AMOUNT_COLUMN_PADDING), 255
        )

    footnote_row = data_max_row + 2
    ws.merge_cells(start_row=footnote_row, start_column=1, end_row=footnote_row, end_column=data_max_col)
    footnote_cell = ws.cell(row=footnote_row, column=1)
    footnote_cell.value = MONTH_DR_CR_FOOTNOTE
    footnote_cell.font = FONT_FOOTNOTE
    footnote_cell.alignment = ALIGN_LEFT

    chart_data_end_row = data_max_row
    if chart_data_end_row >= 3:
        last_label = str(ws.cell(row=chart_data_end_row, column=1).value or "").strip().lower()
        if last_label == "total":
            chart_data_end_row -= 1

    if chart_data_end_row >= 3:
        _add_month_dr_cr_chart_image(ws, chart_data_end_row, footnote_row)


def _apply_repeat_group_colors(workbook, sheet_name: str, amount_column: str) -> None:
    if sheet_name not in workbook.sheetnames:
        return

    ws = workbook[sheet_name]
    if ws.max_row <= 1:
        return

    header_to_col: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        value = ws.cell(row=1, column=col_idx).value
        header_to_col[str(value).strip()] = col_idx

    amount_col_idx = header_to_col.get(amount_column)
    if amount_col_idx is None:
        return

    color_map: dict[str, PatternFill] = {}
    color_index = 0

    for row_idx in range(2, ws.max_row + 1):
        raw_value = ws.cell(row=row_idx, column=amount_col_idx).value
        if raw_value in (None, ""):
            continue

        rounded_value = _round_money_for_excel(raw_value)
        if rounded_value is not None:
            key = str(rounded_value)
        else:
            key = str(raw_value)

        if key not in color_map:
            color_map[key] = REPEAT_GROUP_FILLS[color_index % len(REPEAT_GROUP_FILLS)]
            color_index += 1

        fill = color_map[key]
        for col_idx in range(1, ws.max_column + 1):
            ws.cell(row=row_idx, column=col_idx).fill = fill


def _apply_pdf_review_highlight(ws, header_to_col: dict[str, int]) -> None:
    """Summarize findings; reserve the red highlight for confirmed changes."""
    status_col = header_to_col["status"]
    candidates = [
        (str(ws.cell(row, status_col).value or "").strip().upper(), row)
        for row in range(PDF_STATUS_TABLE_START_ROW + 1, ws.max_row + 1)
        if str(ws.cell(row, status_col).value or "").strip().upper() in PDF_STATUS_RANK
    ]
    severity = max((status for status, _ in candidates), key=PDF_STATUS_RANK.get, default="WARNING")
    worst_rows = [row for status, row in candidates if status == severity]
    check_col = header_to_col.get("check")
    overall_rows = [row for row in worst_rows if check_col and
                    ws.cell(row, check_col).value == "Overall PDF modification status"]
    selected = next(iter(overall_rows or worst_rows), None)
    result_col, pdf_col = header_to_col.get("result"), header_to_col.get("pdf")
    result = ws.cell(selected, result_col).value if selected and result_col else None
    message = str(result or "Review required; see checks below.")
    pdfs = {ws.cell(row, pdf_col).value for _, row in candidates if pdf_col and ws.cell(row, pdf_col).value}
    if len(pdfs) > 1 and selected and pdf_col:
        message = f"Most severe finding across all PDFs. {ws.cell(selected, pdf_col).value}: {message}"

    review_row = PDF_STATUS_TABLE_START_ROW - 2
    ws.cell(review_row, 1, "PDF review:")
    ws.cell(review_row, 2, _sanitize_excel_value(f"{severity}: {message}"))
    ws.merge_cells(start_row=review_row, start_column=2, end_row=review_row, end_column=5)

    if severity != "FAIL":
        return

    for row in (1, review_row):
        for col in range(1, 6):
            cell = ws.cell(row, col)
            cell.fill = PatternFill(fill_type="solid", fgColor="C00000")
            cell.font = Font(name="Aptos", size=16 if row == 1 else 12, bold=True, color="FFFFFF")
            cell.alignment = ALIGN_LEFT


def _apply_pdf_status_style(
    workbook,
    sheet_name: str,
    account_summary_rows: list[tuple[str, str]],
) -> None:
    if sheet_name not in workbook.sheetnames:
        return

    ws = workbook[sheet_name]

    for row_idx, (label, value) in enumerate(account_summary_rows, start=1):
        label_cell = ws.cell(row=row_idx, column=1)
        value_cell = ws.cell(row=row_idx, column=2)
        label_cell.value = _sanitize_excel_value(f"{label}:")
        value_cell.value = _sanitize_excel_value(value)
        if label == "Account Number":
            value_cell.number_format = "@"
        label_cell.font = FONT_HEADER
        value_cell.font = FONT_NORMAL
        label_cell.alignment = ALIGN_LEFT
        value_cell.alignment = ALIGN_LEFT
        for cell in (label_cell, value_cell):
            cell.border = THIN_BORDER
            cell.fill = PatternFill(fill_type="solid", fgColor="F2F2F2")

    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=5)
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=5)
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=5)
    ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=5)
    ws.merge_cells(start_row=5, start_column=2, end_row=5, end_column=5)

    ws.freeze_panes = f"A{PDF_STATUS_TABLE_START_ROW + 1}"
    if ws.max_row >= PDF_STATUS_TABLE_START_ROW:
        ws.auto_filter.ref = (
            f"A{PDF_STATUS_TABLE_START_ROW}:"
            f"{ws.cell(row=PDF_STATUS_TABLE_START_ROW, column=ws.max_column).coordinate[:-1]}{ws.max_row}"
        )

    table_header_row = PDF_STATUS_TABLE_START_ROW
    header_to_col: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        header_cell = ws.cell(row=table_header_row, column=col_idx)
        value = header_cell.value
        header_to_col[_normalize_header(value)] = col_idx
        header_cell.font = FONT_HEADER
        header_cell.alignment = ALIGN_CENTER
        header_cell.fill = HEADER_FILL
        header_cell.border = THIN_BORDER

    status_col_idx = header_to_col.get("status")
    if status_col_idx is None:
        return

    align_by_header = {
        "pdf": ALIGN_LEFT,
        "check": ALIGN_LEFT,
        "status": ALIGN_CENTER,
        "result": ALIGN_LEFT,
        "details": ALIGN_LEFT,
    }
    for row_idx in range(table_header_row + 1, ws.max_row + 1):
        row_fill = ALT_ROW_FILLS[(row_idx - table_header_row - 1) % 2]
        for col_idx in range(1, ws.max_column + 1):
            header = _normalize_header(ws.cell(row=table_header_row, column=col_idx).value)
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = FONT_NORMAL
            cell.alignment = align_by_header.get(header, ALIGN_LEFT)
            cell.fill = row_fill
            cell.border = THIN_BORDER

        cell = ws.cell(row=row_idx, column=status_col_idx)
        status = str(cell.value or "").strip().upper()
        fill = PDF_STATUS_FILLS.get(status)
        if fill is None:
            continue
        cell.fill = fill
        cell.font = Font(name="Aptos", size=10, bold=True)
        cell.alignment = ALIGN_CENTER

    _apply_pdf_review_highlight(ws, header_to_col)

    table_widths: dict[int, float] = {}
    for col_idx in range(1, ws.max_column + 1):
        column_letter = ws.cell(row=table_header_row, column=col_idx).column_letter
        max_length = max(
            (len(line)
             for row_idx in range(table_header_row, ws.max_row + 1)
             for line in str(ws.cell(row=row_idx, column=col_idx).value or "").split("\n")),
            default=0,
        )
        width = min(max(max_length + 4, 10), PDF_STATUS_MAX_COLUMN_WIDTH)
        ws.column_dimensions[column_letter].width = width
        table_widths[col_idx] = width

    for row_idx in range(1, ws.max_row + 1):
        if row_idx <= PDF_STATUS_TABLE_START_ROW - 2:
            label = str(ws.cell(row=row_idx, column=1).value or "")
            value = str(ws.cell(row=row_idx, column=2).value or "")
            font_size = max(ws.cell(row_idx, col).font.sz or 10 for col in (1, 2))
            label_width = max(1, int((table_widths.get(1, 10) - 3) * 10 / font_size))
            value_width = max(1, int((sum(table_widths.get(col, 10) for col in range(2, 6)) - 3) * 10 / font_size))
            label_lines = sum(max(1, (len(line) + label_width - 1) // label_width) for line in label.split("\n"))
            value_lines = sum(max(1, (len(line) + value_width - 1) // value_width) for line in value.split("\n"))
            ws.row_dimensions[row_idx].height = min(409, max(30, font_size * 1.4 * max(label_lines, value_lines) + 12))
        else:
            line_counts = []
            for col, width in table_widths.items():
                text = str(ws.cell(row=row_idx, column=col).value or "")
                usable_width = max(1, int(width) - 3)
                line_counts.append(
                    sum(max(1, (len(line) + usable_width - 1) // usable_width) for line in text.split("\n"))
                )
            ws.row_dimensions[row_idx].height = min(409, max(36, 14 * max(line_counts) + 8))


def _force_leading_equals_to_text(workbook) -> None:
    force_leading_equals_to_text(workbook)


def _next_final_path(output_dir: Path, pdf_stem: str) -> Path:
    return append_timestamp_if_exists(output_dir / f"{pdf_stem}.xlsx")


def build_final_workbook(
    statement_df: pd.DataFrame,
    rules_path: Path,
    output_dir: Path,
    pdf_stem: str,
    logger,
    source_pdf_paths: list[Path] | None = None,
    source_pdf_passwords: list[str | None] | None = None,
    source_audits: list[dict[str, object]] | None = None,
    include_source: bool = True,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    statement_df = _ensure_columns(statement_df)

    pdf_account_summary_rows = _build_pdf_account_summary_rows(source_pdf_paths, source_pdf_passwords)
    customer_name = dict(pdf_account_summary_rows).get("Customer Name", "")
    rules = _load_rules(rules_path, logger)
    rule_sheets = _build_rule_sheets(statement_df, rules, logger)

    pdf_status_df = _build_pdf_status_sheet(source_pdf_paths, source_pdf_passwords, source_audits)
    return_reject_df = _build_return_reject_sheet(statement_df)
    cheque_df = _build_cheque_sheet(statement_df)
    repeat_credit_df = _build_repeat_sheet(statement_df, "Credit")
    repeat_debit_df = _build_repeat_sheet(statement_df, "Debit")
    top30_debit_df = _build_top_sheet(statement_df, "Debit", top_n=30)
    top30_credit_df = _build_top_sheet(statement_df, "Credit", top_n=30)
    month_dr_cr_df = _build_month_dr_cr_sheet(statement_df)

    planned_sheets: list[tuple[str, pd.DataFrame]] = [
        (PDF_STATUS_SHEET_NAME, pdf_status_df),
        ("Statement", statement_df),
        ("Ret/Rej", return_reject_df),
    ]
    planned_sheets.extend(rule_sheets)
    planned_sheets.extend(
        [
            ("Cheque_Transactions", cheque_df),
            ("Repeat_Credit_Amount", repeat_credit_df),
            ("Repeat_Debit_Amount", repeat_debit_df),
            ("Top30_Debit", top30_debit_df),
            ("Top30_Credit", top30_credit_df),
            ("month_dr_cr", month_dr_cr_df),
        ]
    )

    final_path = _next_final_path(output_dir, pdf_stem)

    used_names: set[str] = set()
    normalized_sheet_names: dict[str, str] = {}
    rule_schedule_sheets = []

    with pd.ExcelWriter(final_path, engine="openpyxl") as writer:
        for requested_name, frame in planned_sheets:
            safe_name = _unique_sheet_name(requested_name, used_names)
            normalized_sheet_names[requested_name] = safe_name
            if any(frame is rule_frame for _, rule_frame in rule_sheets):
                schedule = build_rule_schedule(frame, rules, requested_name, align_rows=True)
                if not schedule.empty:
                    frame = order_rule_transactions(frame)
                rule_schedule_sheets.append((safe_name, schedule))
            display_frame = _exclude_source_column(
                _sanitize_excel_frame(_exclude_final_columns(frame)),
                include_source,
            )
            if requested_name == PDF_STATUS_SHEET_NAME:
                display_frame.to_excel(
                    writer,
                    sheet_name=safe_name,
                    index=False,
                    startrow=PDF_STATUS_TABLE_START_ROW - 1,
                )
            elif requested_name == "month_dr_cr":
                display_frame.to_excel(writer, sheet_name=safe_name, index=False)
            else:
                _exclude_source_column(
                    _sanitize_excel_frame(
                        _exclude_final_columns(_ensure_columns(frame))
                    ),
                    include_source,
                ).to_excel(writer, sheet_name=safe_name, index=False)

    workbook = load_workbook(final_path)
    _apply_base_style(workbook)
    _apply_repeat_group_colors(
        workbook,
        normalized_sheet_names.get("Repeat_Credit_Amount", "Repeat_Credit_Amount"),
        "Credit",
    )
    _apply_repeat_group_colors(
        workbook,
        normalized_sheet_names.get("Repeat_Debit_Amount", "Repeat_Debit_Amount"),
        "Debit",
    )
    _apply_month_dr_cr_style(
        workbook,
        normalized_sheet_names.get("month_dr_cr", "month_dr_cr"),
        customer_name,
    )
    _apply_pdf_status_style(
        workbook,
        normalized_sheet_names.get(PDF_STATUS_SHEET_NAME, PDF_STATUS_SHEET_NAME),
        pdf_account_summary_rows,
    )
    _apply_final_layout(workbook)
    cheque_ranges = {}
    for sheet_name, schedule in rule_schedule_sheets:
        cheque_range = add_rule_schedule(workbook[sheet_name], schedule)
        if cheque_range:
            cheque_ranges[workbook.sheetnames.index(sheet_name) + 1] = cheque_range
    _force_leading_equals_to_text(workbook)
    workbook.save(final_path)
    workbook.close()
    suppress_number_as_text_warnings(final_path, cheque_ranges)

    logger.info("Final workbook created: %s", final_path)
    return final_path
