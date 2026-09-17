"""Render the monthly debit/credit chart embedded in the workbook."""

from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage, ImageDraw, ImageFont

MONTH_DR_CR_CHART_IMAGE_SIZE = (1120, 520)
MONTH_DR_CR_DATA_LABEL_FONT_SIZE = 13


def _format_month_dr_cr_chart_label(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""

    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return ""

    absolute = abs(numeric)
    if absolute >= 10000000:
        return f"{numeric / 10000000:.2f} Cr"
    if absolute >= 100000:
        return f"{numeric / 100000:.1f} L"
    if absolute >= 1000:
        return f"{numeric / 1000:.1f} k"
    return f"{numeric:.1f}"


def _load_chart_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    font_names = ["arialbd.ttf", "calibrib.ttf", "arial.ttf", "calibri.ttf"] if bold else [
        "arial.ttf",
        "calibri.ttf",
    ]
    for font_name in font_names:
        try:
            return ImageFont.truetype(font_name, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def _draw_dotted_horizontal_line(draw: ImageDraw.ImageDraw, x1: int, x2: int, y: int, fill: str) -> None:
    dash_length = 4
    gap_length = 6
    x = x1
    while x < x2:
        draw.line((x, y, min(x + dash_length, x2), y), fill=fill, width=1)
        x += dash_length + gap_length


def _format_month_dr_cr_axis_label(value: float) -> str:
    return f"\u20B9{_format_month_dr_cr_chart_label(value)}"


def _nice_axis_step(max_value: float, tick_count: int = 8) -> float:
    if max_value <= 0:
        return 1.0

    rough_step = max_value / max(tick_count, 1)
    exponent = int(f"{rough_step:e}".split("e")[1])
    magnitude = 10 ** exponent
    fraction = rough_step / magnitude

    if fraction <= 1:
        nice_fraction = 1
    elif fraction <= 2:
        nice_fraction = 2
    elif fraction <= 5:
        nice_fraction = 5
    else:
        nice_fraction = 10

    return nice_fraction * magnitude


def _draw_rotated_text(
    image: PILImage.Image,
    text: str,
    position: tuple[int, int],
    font,
    fill: str,
    angle: float,
) -> None:
    if not text:
        return

    measure_draw = ImageDraw.Draw(PILImage.new("RGBA", (1, 1), (0, 0, 0, 0)))
    bbox = measure_draw.textbbox((0, 0), text, font=font)
    text_width = max(1, bbox[2] - bbox[0])
    text_height = max(1, bbox[3] - bbox[1])
    text_layer = PILImage.new("RGBA", (text_width + 8, text_height + 8), (0, 0, 0, 0))
    text_draw = ImageDraw.Draw(text_layer)
    text_draw.text((4, 4), text, font=font, fill=fill)
    rotated = text_layer.rotate(angle, expand=True, resample=PILImage.Resampling.BICUBIC)
    image.alpha_composite(rotated, dest=position)


def _add_month_dr_cr_chart_image(ws, chart_data_end_row: int, footnote_row: int) -> None:
    month_values: list[tuple[str, float, float]] = []
    for row_idx in range(2, chart_data_end_row + 1):
        month_label = str(ws.cell(row=row_idx, column=1).value or "").strip()
        if not month_label:
            continue
        try:
            debit = float(ws.cell(row=row_idx, column=2).value or 0)
            credit = float(ws.cell(row=row_idx, column=3).value or 0)
        except (TypeError, ValueError):
            continue
        month_values.append((month_label, debit, credit))

    if not month_values:
        return

    width, height = MONTH_DR_CR_CHART_IMAGE_SIZE
    image = PILImage.new("RGBA", (width, height), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)

    panel_margin = 12
    panel_radius = 26
    shadow_offset = 8
    shadow_color = (0, 0, 0, 48)
    panel_fill = "#272624"
    panel_border = "#4A4744"
    grid_color = "#403D39"
    axis_text_color = "#8F8B86"
    label_text_color = "#BFBAB3"
    value_text_color = "#D9D5CF"
    bar_colors = {"Dr": "#355C91", "Cr": "#B0572D"}

    draw.rounded_rectangle(
        (
            panel_margin + shadow_offset,
            panel_margin + shadow_offset,
            width - panel_margin + shadow_offset,
            height - panel_margin + shadow_offset,
        ),
        radius=panel_radius,
        fill=shadow_color,
    )
    draw.rounded_rectangle(
        (panel_margin, panel_margin, width - panel_margin, height - panel_margin),
        radius=panel_radius,
        fill=panel_fill,
        outline=panel_border,
        width=2,
    )

    font_regular = _load_chart_font(12)
    font_label = _load_chart_font(MONTH_DR_CR_DATA_LABEL_FONT_SIZE)
    font_bold = _load_chart_font(12, bold=True)

    left_margin = panel_margin + 78
    right_margin = panel_margin + 34
    top_margin = panel_margin + 92
    bottom_margin = panel_margin + 88
    plot_left = left_margin
    plot_top = top_margin
    plot_right = width - right_margin
    plot_bottom = height - bottom_margin
    plot_width = plot_right - plot_left
    plot_height = plot_bottom - plot_top

    max_value = max(max(debit, credit) for _, debit, credit in month_values)
    if max_value <= 0:
        max_value = 1.0
    axis_step = _nice_axis_step(max_value, tick_count=8)
    axis_max = axis_step * max(1, int((max_value + axis_step - 1) // axis_step))

    def value_to_y(value: float) -> int:
        scaled = value / axis_max
        return round(plot_bottom - (plot_height * scaled))

    tick_value = 0.0
    while tick_value <= axis_max + (axis_step / 2):
        y = value_to_y(tick_value)
        _draw_dotted_horizontal_line(draw, plot_left, plot_right, y, grid_color)
        tick_value += axis_step

    group_width = plot_width / max(len(month_values), 1)
    bar_width = max(16, min(34, int(group_width * 0.30)))
    series_gap = max(8, int(bar_width * 0.28))

    label_offsets = {"Dr": -12, "Cr": 10}

    for index, (month_label, debit, credit) in enumerate(month_values):
        group_center = plot_left + group_width * (index + 0.5)

        dr_left = round(group_center - series_gap / 2 - bar_width)
        dr_right = dr_left + bar_width
        dr_top = value_to_y(debit)
        draw.rounded_rectangle((dr_left, dr_top, dr_right, plot_bottom), radius=7, fill=bar_colors["Dr"])

        cr_left = round(group_center + series_gap / 2)
        cr_right = cr_left + bar_width
        cr_top = value_to_y(credit)
        draw.rounded_rectangle((cr_left, cr_top, cr_right, plot_bottom), radius=7, fill=bar_colors["Cr"])

        dr_label = _format_month_dr_cr_chart_label(debit)
        dr_bbox = draw.textbbox((0, 0), dr_label, font=font_label)
        _draw_rotated_text(
            image,
            dr_label,
            (
                int((dr_left + dr_right) // 2 + label_offsets["Dr"] - ((dr_bbox[2] - dr_bbox[0]) * 0.14)),
                int(max(plot_top - 8, dr_top - 42)),
            ),
            font_label,
            label_text_color,
            64,
        )

        cr_label = _format_month_dr_cr_chart_label(credit)
        cr_bbox = draw.textbbox((0, 0), cr_label, font=font_label)
        _draw_rotated_text(
            image,
            cr_label,
            (
                int((cr_left + cr_right) // 2 + label_offsets["Cr"] - ((cr_bbox[2] - cr_bbox[0]) * 0.12)),
                int(max(plot_top - 8, cr_top - 42)),
            ),
            font_label,
            label_text_color,
            64,
        )

        month_bbox = draw.textbbox((0, 0), month_label, font=font_regular)
        _draw_rotated_text(
            image,
            month_label,
            (
                int(round(group_center) - ((month_bbox[2] - month_bbox[0]) * 0.55)),
                int(plot_bottom + 4),
            ),
            font_regular,
            axis_text_color,
            45,
        )

    legend_x = panel_margin + 28
    legend_y = panel_margin + 28
    legend_cursor = legend_x
    for legend_label, legend_color in (("Debit (Dr)", bar_colors["Dr"]), ("Credit (Cr)", bar_colors["Cr"])):
        draw.rounded_rectangle((legend_cursor, legend_y + 4, legend_cursor + 16, legend_y + 20), radius=4, fill=legend_color)
        draw.text((legend_cursor + 22, legend_y), legend_label, font=font_bold, fill=value_text_color)
        label_bbox = draw.textbbox((0, 0), legend_label, font=font_bold)
        legend_cursor += 22 + (label_bbox[2] - label_bbox[0]) + 28

    image_bytes = BytesIO()
    image.save(image_bytes, format="PNG")
    image_bytes.seek(0)

    chart_image = XLImage(image_bytes)
    chart_image.width = width
    chart_image.height = height
    chart_row = footnote_row + 5
    ws.add_image(chart_image, f"A{chart_row}")
