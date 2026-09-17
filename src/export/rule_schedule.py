"""Append the reference-style schedule beside a rule sheet's original table."""

from copy import copy

import pandas as pd
from openpyxl.styles import Alignment, Border, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.export.excel_safety import sanitize_excel_value
from src.transform.rule_schedule import SCHEDULE_COLUMNS

SCHEDULE_HEADERS = [
    "Due No", "Name", "Actual Date", "Day", "Freq", "Cheque No",
    "Debit", "Credit", "Paid Date", "Other Name",
]


def add_rule_schedule(ws, schedule: pd.DataFrame) -> str | None:
    """Add the schedule and return its cheque range for Excel error suppression."""
    if schedule.empty:
        return

    source_last_col = ws.max_column
    start_col = source_last_col + 2
    end_col = start_col + len(SCHEDULE_COLUMNS) - 1
    title_row, header_row, first_row = 1, 2, 3
    last_row = first_row + len(schedule) - 1
    font = copy(ws.cell(2, 1).font)
    font.color = "000000"
    credit_font = copy(font)
    credit_font.color = "FF0000"
    header_font = copy(ws.cell(1, 1).font)
    title_font = copy(header_font)
    title_font.color = "000000"
    header_font.color = "FFFFFF"
    line = Side(style="thin", color="000000")
    no_line = Side(style=None)
    title_fill = PatternFill("solid", fgColor="92CDDC")
    source_fill = PatternFill("solid", fgColor="B7DEE8")
    header_fill = PatternFill("solid", fgColor="31869B")
    body_fills = [PatternFill("solid", fgColor="EAF3FB"), PatternFill("solid", fgColor="F8F1E5")]
    center = Alignment(horizontal="center", vertical="center")
    amount_headers = {"Debit", "Credit", "Balance"}
    source_amount_cols = {c for c in range(1, source_last_col + 1)
                          if ws.cell(1, c).value in amount_headers}

    # Both tables share their title, header and transaction rows.
    ws.insert_rows(1)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(end_col)}{last_row}"
    ws.row_dimensions[title_row].height = 43
    ws.row_dimensions[header_row].height = 19.5
    for left, right, fill in ((1, source_last_col, source_fill), (start_col, end_col, title_fill)):
        ws.merge_cells(start_row=title_row, start_column=left, end_row=title_row, end_column=right)
        for col in range(left, right + 1):
            cell = ws.cell(title_row, col)
            cell.fill = fill
            cell.border = Border(left=line if col == left else no_line,
                                 right=line if col == right else no_line,
                                 top=line, bottom=line)
        title = ws.cell(title_row, left, sanitize_excel_value(ws.title))
        title.font = title_font
        title.alignment = center

    for row in ws.iter_rows(min_row=header_row, max_row=last_row, min_col=1, max_col=source_last_col):
        for cell in row:
            is_header = cell.row == header_row
            cell.font = title_font if is_header else font
            cell.fill = source_fill if is_header else body_fills[(cell.row - first_row) % 2]
            if not is_header and cell.column in source_amount_cols:
                # Match the reference's locale-aware, whole-number display.
                cell.number_format = "###,##0"
            cell.border = Border(
                left=line if cell.column == 1 else no_line,
                right=line if cell.column == source_last_col else no_line,
                bottom=line if is_header or cell.row == last_row else no_line,
            )

    # Reference: one group through FREQ, then mode, DR, CR and manual fields.
    group_right_edges = {4, 5, 6, 7, 9}
    for offset, label in enumerate(SCHEDULE_HEADERS):
        cell = ws.cell(header_row, start_col + offset, label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = Border(
            left=line if offset == 0 else no_line,
            right=line if offset in group_right_edges else no_line,
            top=line, bottom=line,
        )

    source_headers = {ws.cell(header_row, c).value: c for c in range(1, start_col - 1)}
    for row_index, values in enumerate(schedule.itertuples(index=False, name=None), first_row):
        is_credit = pd.notna(values[7]) and values[7] != 0
        ws.row_dimensions[row_index].height = 19
        for offset, value in enumerate(values):
            if pd.isna(value):
                value = None
            cell = ws.cell(row_index, start_col + offset, sanitize_excel_value(value))
            cell.font = credit_font if is_credit else font
            cell.fill = body_fills[(row_index - first_row) % 2]
            cell.alignment = Alignment(
                horizontal="right" if offset in {6, 7} else "left" if offset in {1, 2, 3, 8, 9} else "center",
                vertical="center",
            )
            cell.border = Border(
                left=line if offset == 0 else no_line,
                right=line if offset in group_right_edges else no_line,
                bottom=line if row_index == last_row else no_line,
            )
            if offset in {2, 8}:
                cell.number_format = "yyyy-mm-dd"
            elif offset in {6, 7}:
                source_col = source_headers.get("Debit" if offset == 6 else "Credit")
                cell.number_format = ws.cell(first_row, source_col).number_format if source_col else "#,##,##0"
            elif offset == 5:
                cell.number_format = "@"
            elif offset in {0, 4}:
                cell.number_format = "0"

    ws.row_dimensions[last_row + 1].height = 19
    for offset, label in ((6, "DR"), (7, "CR")):
        cell = ws.cell(last_row + 1, start_col + offset, float(schedule[label].sum()))
        cell.font = font
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.number_format = ws.cell(first_row, start_col + offset).number_format
        cell.border = Border(left=line, right=line, top=line, bottom=line)

    for offset, label in enumerate(SCHEDULE_HEADERS):
        letter = get_column_letter(start_col + offset)
        source_label = {2: "Date", 6: "Debit", 7: "Credit", 8: "Date"}.get(offset)
        if source_label in source_headers:
            width = ws.column_dimensions[get_column_letter(source_headers[source_label])].width
        elif offset == 3:
            width = 10
        elif offset in {6, 7}:
            width = 16
        else:
            lengths = [len(str(v)) for v in schedule.iloc[:, offset] if pd.notna(v)] if offset in {1, 3, 5} else []
            padding = 5 if offset == 1 else 2
            width = min(max([10, len(label) + padding] + [length + padding for length in lengths]), 255)
        ws.column_dimensions[letter].width = width
    ws.column_dimensions[get_column_letter(start_col - 1)].width = 15
    cheque_letter = get_column_letter(start_col + 5)
    return f"{cheque_letter}{first_row}:{cheque_letter}{last_row}"
