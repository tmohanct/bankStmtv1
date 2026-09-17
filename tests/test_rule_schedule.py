from __future__ import annotations

import logging
import tempfile
import unittest
from copy import copy
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook

from src.export.final_excel_builder import build_final_workbook
from src.export.rule_schedule import SCHEDULE_HEADERS
from src.transform.analysis import _build_rule_sheets, _ensure_columns, _load_rules
from src.transform.rule_schedule import SCHEDULE_COLUMNS, build_rule_schedule, transaction_mode
from src.utils.text_utils import clean_detail


RULES = [
    {"category": "FIN", "name": "RAMASAMY", "name_clean": "RAMASAMY", "sheet_name": "RAMASAMY"},
    {"category": "FIN", "name": "KANCHANA", "name_clean": "KANCHANA", "sheet_name": "ramasamy"},
]


def statement():
    entries = [
        ("20/03/2026", "CHQ 000607 KANCHANA", "000607", 85000.25, 0),
        ("12/03/2026", "RTGS/KANCHANA", "987654321", 0, 800000),
        ("21/05/2026", "IMPS/123/RAMASAMY", "123456789", 85000, 0),
        ("13/03/2026", "NEFT/RAMASAMY", "", 0, 800000),
        ("26/03/2026", "CHQ 608 RAMASAMY", "608", 85000, 0),
        ("26/03/2026", "UPI/RAMASAMY/KANCHANA", "", 25, 0),
    ]
    return _ensure_columns(pd.DataFrame([
        {"Sno": i, "Date": date, "Details": detail, "Detail_Clean": clean_detail(detail),
         "Cheque No": cheque, "Debit": debit, "Credit": credit, "Balance": 1000000,
         "Source": "sample.pdf"}
        for i, (date, detail, cheque, debit, credit) in enumerate(entries, 1)
    ]))


class RuleScheduleTests(unittest.TestCase):
    def test_two_column_rules_merge_matches_and_preserve_schedule_names(self):
        with tempfile.TemporaryDirectory(prefix="search_string_rules_") as tmp:
            rules_path = Path(tmp) / "Rules.xlsx"
            for search_header in ("searchString", " SEARCHSTRING "):
                with self.subTest(search_header=search_header):
                    pd.DataFrame([
                        {search_header: "RAMASAMY", "SheetName": "RAMASAMY"},
                        {search_header: "KANCHANA", "SheetName": "RAMASAMY"},
                    ]).to_excel(rules_path, index=False)
                    rules = _load_rules(rules_path, logging.getLogger(__name__))
                    sheets = _build_rule_sheets(statement(), rules, logging.getLogger(__name__))
                    self.assertEqual([name for name, _ in sheets], ["RAMASAMY"])
                    self.assertEqual(sheets[0][1]["Sno"].tolist(), [1, 2, 3, 4, 5, 6])
                    schedule = build_rule_schedule(sheets[0][1], rules, "RAMASAMY")
                    self.assertEqual(schedule["NAME"].tolist(), [
                        "KANCHANA", "RAMASAMY", "KANCHANA", "RAMASAMY", "RAMASAMY", "RAMASAMY",
                    ])

    def test_date_order_names_gaps_debit_sequence_and_modes(self):
        frame = statement()
        original = frame.copy(deep=True)
        schedule = build_rule_schedule(frame, RULES, "RAMASAMY")
        self.assertEqual(list(schedule.columns), SCHEDULE_COLUMNS)
        self.assertEqual(schedule["NAME"].tolist(), ["KANCHANA", "RAMASAMY", "KANCHANA", "RAMASAMY", "RAMASAMY", "RAMASAMY"])
        self.assertTrue(schedule["FREQ"].iloc[:2].isna().all())
        self.assertEqual(schedule["FREQ"].iloc[2], 8)
        self.assertEqual(schedule["FREQ"].iloc[3:].tolist(), [6, 0, 56])
        self.assertTrue(schedule["DUE NO"].iloc[:2].isna().all())
        self.assertEqual(schedule["DUE NO"].iloc[2:].tolist(), [1, 2, 3, 4])
        self.assertEqual(schedule["CHEQUE NO"].tolist(), ["RTGS", "NEFT", "000607", "608", "UPI", "IMPS"])
        self.assertEqual(schedule["DAY"].iloc[0], "Thursday")
        self.assertEqual(schedule["DR"].sum(), 255025.25)
        self.assertEqual(schedule["CR"].sum(), 1600000)
        self.assertTrue(schedule[["PAID DATE", "OTHER NAME"]].isna().all().all())
        pd.testing.assert_frame_equal(frame, original)

    def test_credit_between_debits_does_not_reset_frequency(self):
        frame = statement()
        frame.loc[1, "Date"] = "23/03/2026"
        schedule = build_rule_schedule(frame, RULES, "RAMASAMY")
        self.assertTrue(schedule.loc[schedule["CR"].notna(), "FREQ"].isna().all())
        debits = schedule[schedule["DR"].notna()]
        self.assertEqual(debits.iloc[0]["FREQ"], 7)
        self.assertEqual(debits["FREQ"].iloc[1:].tolist(), [6, 0, 56])

    def test_duplicate_matches_remain_one_row_and_first_rule_wins(self):
        sheets = _build_rule_sheets(statement(), RULES, logging.getLogger(__name__))
        self.assertEqual(len(sheets), 1)
        schedule = build_rule_schedule(sheets[0][1], RULES, sheets[0][0])
        self.assertEqual(len(schedule), 6)
        self.assertEqual(schedule.iloc[4]["NAME"], "RAMASAMY")

    def test_first_due_without_preceding_credit_stays_unavailable(self):
        frame = statement()
        frame.loc[frame["Credit"] > 0, "Date"] = "24/03/2026"
        schedule = build_rule_schedule(frame, RULES, "RAMASAMY")
        debits = schedule[schedule["DR"].notna()]
        self.assertTrue(pd.isna(debits.iloc[0]["FREQ"]))
        self.assertEqual(debits["FREQ"].iloc[1:].tolist(), [6, 0, 56])

    def test_aligned_schedule_reserves_amount_only_rows(self):
        frame = statement()
        frame.loc[0, ["Details", "Detail_Clean"]] = "UNRELATED"
        schedule = build_rule_schedule(frame, RULES, "RAMASAMY", align_rows=True)
        self.assertEqual(len(schedule), len(frame))
        self.assertTrue(schedule.iloc[2].isna().all())
        self.assertEqual(schedule.iloc[3]["DUE NO"], 1)
        self.assertEqual(schedule.iloc[3]["FREQ"], 14)

    def test_missing_dates_sort_last_without_invented_days_or_gaps(self):
        frame = statement()
        frame.loc[0, "Date"] = "unreadable"
        frame.loc[1, "Date"] = None
        schedule = build_rule_schedule(frame, RULES, "RAMASAMY")
        self.assertEqual(schedule.iloc[-2]["ACTUAL DATE"], "unreadable")
        self.assertTrue(schedule.iloc[-2:][["DAY", "FREQ"]].isna().all().all())

    def test_amount_only_and_unmatched_rules_have_no_schedule(self):
        rules = [{"category": "AMT", "name": "85000", "amount_value": 85000, "sheet_name": "Amount"}]
        self.assertTrue(build_rule_schedule(statement(), rules, "Amount").empty)
        self.assertTrue(build_rule_schedule(statement(), RULES, "Other").empty)
        self.assertTrue(build_rule_schedule(statement().iloc[:0], RULES, "RAMASAMY").empty)

    def test_mode_does_not_guess_from_names_or_transfer_reference(self):
        for detail, expected in [("IMPS123/RAMASAMY", "IMPS"), ("upi/123", "UPI"),
                                 ("BY ONL 0000IMPSICI607115003462:RAMASAMY", "IMPS"),
                                 ("TO ONL IMPSCUB614110789302:RAMASAMY", "IMPS"),
                                 ("Paid to KUPINDER", ""), ("CHQ CLEARING", "CHEQUE"),
                                 ("RTGS/123", "RTGS")]:
            with self.subTest(detail=detail):
                self.assertEqual(transaction_mode({"Details": detail, "Cheque No": ""}), expected)

    def test_mixed_amount_destination_excludes_rows_without_matching_name(self):
        frame = statement()
        frame.loc[0, "Detail_Clean"] = "UNRELATED"
        rules = RULES + [{"category": "AMT", "name": "85000", "amount_value": 85000, "sheet_name": "RAMASAMY"}]
        schedule = build_rule_schedule(frame, rules, "RAMASAMY")
        self.assertEqual(len(schedule), 5)
        self.assertNotIn("000607", schedule["CHEQUE NO"].tolist())

    def test_export_aligns_transactions_and_formats_reference_layout(self):
        with tempfile.TemporaryDirectory(prefix="rule_schedule_") as tmp:
            root = Path(tmp)
            rules_path = root / "Rules.xlsx"
            pd.DataFrame([
                {"Category": "FIN", "subCategory": "RAMASAMY", "SheetName": "RAMASAMY"},
                {"Category": "FIN", "subCategory": "KANCHANA", "SheetName": "ramasamy"},
                {"Category": "AMT", "subCategory": "85000", "SheetName": "Amounts"},
            ]).to_excel(rules_path, index=False)
            for include_source, start_col in [(True, 10), (False, 9)]:
                with self.subTest(include_source=include_source):
                    result = build_final_workbook(statement(), rules_path, root, "sample", logging.getLogger(__name__), include_source=include_source)
                    with result.open("rb") as handle:
                        wb = load_workbook(handle)
                        ws = wb["RAMASAMY"]
                        self.assertEqual(ws["A1"].value, ws.title)
                        self.assertEqual(ws.cell(1, start_col).value, ws.title)
                        self.assertEqual([ws.cell(2, start_col+i).value for i in range(10)], SCHEDULE_HEADERS)
                        self.assertEqual([ws.cell(i, 1).value for i in range(3, 9)], [2, 4, 1, 5, 6, 3])
                        self.assertEqual(ws.auto_filter.ref, "A2:S8" if include_source else "A2:R8")
                        self.assertEqual(ws.freeze_panes, "A3")
                        self.assertEqual(ws.row_dimensions[1].height, 43)
                        self.assertEqual(ws.row_dimensions[2].height, 19.5)
                        self.assertEqual(ws.cell(3, start_col+2).value, datetime(2026, 3, 12))
                        self.assertEqual(ws.cell(5, start_col+4).value, 8)
                        self.assertEqual(ws.cell(5, start_col+5).value, "000607")
                        self.assertEqual(ws.cell(5, start_col+5).number_format, "@")
                        self.assertEqual(ws.cell(5, start_col+6).value, 85000.25)
                        self.assertEqual(ws.cell(5, start_col+6).number_format, "###,##0")
                        self.assertEqual(ws["E5"].number_format, "###,##0")
                        self.assertEqual(ws.cell(9, start_col+6).value, 255025.25)
                        self.assertEqual(ws.cell(9, start_col+7).value, 1600000)
                        self.assertEqual(ws.cell(3, start_col).font.name, ws["A3"].font.name)
                        self.assertEqual(ws.cell(3, start_col).font.sz, ws["A3"].font.sz)
                        self.assertEqual(ws["A1"].fill.fgColor.rgb, "00B7DEE8")
                        self.assertEqual(ws["A2"].fill.fgColor.rgb, "00B7DEE8")
                        self.assertEqual(ws.cell(1, start_col).fill.fgColor.rgb, "0092CDDC")
                        self.assertEqual(ws.cell(2, start_col).fill.fgColor.rgb, "0031869B")
                        self.assertEqual(ws.cell(2, start_col).font.color.rgb, "00FFFFFF")
                        self.assertEqual(ws.cell(3, start_col+4).border.right.style, "thin")
                        self.assertIsNone(ws.cell(3, start_col+1).border.right.style)
                        self.assertIsNone(ws.cell(3, start_col+1).border.bottom.style)
                        self.assertEqual(ws.cell(8, start_col+1).border.bottom.style, "thin")
                        for row in range(3, 9):
                            self.assertEqual(ws.cell(row, 2).value, ws.cell(row, start_col+2).value)
                            self.assertEqual(ws.cell(row, 5).value or 0, ws.cell(row, start_col+6).value or 0)
                            self.assertEqual(ws.cell(row, 6).value or 0, ws.cell(row, start_col+7).value or 0)
                            self.assertEqual(copy(ws.cell(row, 1).fill), copy(ws.cell(row, start_col).fill))
                            self.assertEqual(ws.cell(row, 1).font.color.rgb, "00000000")
                            self.assertIsNone(ws.cell(row, start_col-1).value)
                            self.assertIsNone(ws.cell(row, start_col-1).fill.patternType)
                        for row in (3, 4):
                            self.assertIsNone(ws.cell(row, start_col+4).value)
                            for col in range(start_col, start_col+10):
                                self.assertEqual(ws.cell(row, col).font.color.rgb, "00FF0000")
                        self.assertEqual(copy(ws.cell(5, start_col).font), copy(ws["A3"].font))
                        self.assertEqual(ws.column_dimensions[ws.cell(1, start_col-1).column_letter].width, 15)
                        self.assertEqual(ws.column_dimensions[ws.cell(2, start_col+1).column_letter].width, 13)
                        for offset in (8, 9):
                            self.assertTrue(all(ws.cell(r, start_col+offset).value is None for r in range(3, 9)))
                        self.assertEqual(wb["Amounts"].max_column, 8 if include_source else 7)
                        sheet_index = wb.sheetnames.index(ws.title) + 1
                        with ZipFile(result) as archive:
                            xml = ElementTree.fromstring(archive.read(f"xl/worksheets/sheet{sheet_index}.xml"))
                        errors = xml.findall("{*}ignoredErrors/{*}ignoredError")
                        self.assertEqual(len(errors), 1)
                        cheque_letter = ws.cell(5, start_col+5).column_letter
                        self.assertEqual(errors[0].attrib, {
                            "sqref": f"{cheque_letter}3:{cheque_letter}8", "numberStoredAsText": "1",
                        })
                        wb.close()


if __name__ == "__main__":
    unittest.main()
