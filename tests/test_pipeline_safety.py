from __future__ import annotations

import contextlib
from datetime import datetime
import io
import logging
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch
import zipfile

import pandas as pd
import fitz
from openpyxl import load_workbook

import build_fresh_machine_package as packager
from src import main
from src.export.excel_writer import write_output_excel
from src.export.final_excel_builder import (
    _next_final_path,
)
from src.parsers import detector
from src.parsers.parser_registry import PARSER_REGISTRY
from src.transform.normalize import records_to_dataframe
from src.transform.validate import ReconciliationResult, check_running_balances, reconcile, validate_records

ROOT = Path(__file__).resolve().parents[1]


def record(**changes):
    row = dict(Sno=1, Date="01/01/2026", Details="Payment", Detail_Clean="Payment",
               Debit=100.49, Credit=None, Balance=899.51)
    row.update(changes)
    return row


class PipelineValidationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        history_env = patch.dict(os.environ, {"BANKSTMT_DB_PATH": str(self.root / "history.db")})
        history_env.start()
        self.addCleanup(history_env.stop)
        (self.root / "input").mkdir()
        self.pdf = self.root / "input" / "sample.pdf"
        self.pdf.write_bytes(b"fixture path only; parser mocked")

    def run_pipeline(self, records, result):
        with patch.object(main, "PROJECT_ROOT", self.root), patch.object(main, "prepare_pdf_for_reading", return_value=self.pdf), patch.dict(main.PARSERS, {"axis": Mock(return_value=records)}), patch.object(main, "reconcile", return_value=result), patch.object(main, "write_output_excel") as intermediate, patch.object(main, "build_final_workbook") as final, contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            code = main.main(["--bank", "axis", "--pdf", str(self.pdf)])
        return code, intermediate, final

    def test_empty_parse_fails_before_any_workbook(self):
        code, intermediate, final = self.run_pipeline([], ReconciliationResult("unavailable"))
        self.assertEqual(code, 1)
        intermediate.assert_not_called()
        final.assert_not_called()

    def test_summary_mismatch_fails_before_any_workbook(self):
        code, intermediate, final = self.run_pipeline([record()], ReconciliationResult("failed", ("total_debit mismatch",)))
        self.assertEqual(code, 1)
        intermediate.assert_not_called()
        final.assert_not_called()

    def test_unavailable_summary_does_not_claim_a_mismatch(self):
        code, intermediate, final = self.run_pipeline([record()], ReconciliationResult("unavailable"))
        self.assertEqual(code, 0)
        intermediate.assert_called_once()
        final.assert_called_once()

    def test_embedded_filename_password_is_not_logged(self):
        self.pdf = self.pdf.with_name('sample$secret-value.pdf')
        self.pdf.write_bytes(b'fixture path only; parser mocked')
        code, _, _ = self.run_pipeline([record()], ReconciliationResult('unavailable'))
        self.assertEqual(code, 0)
        logs = list((self.root / 'src' / 'logs').glob('*.log'))
        self.assertEqual(len(logs), 1)
        self.assertNotIn('secret-value', logs[0].read_text(encoding='utf-8'))

    def test_invalid_records_are_rejected(self):
        for changes in ({"Date": "31/02/2026"}, {"Debit": "invalid"}, {"Balance": float("inf")}, {"Credit": 10}, {"Debit": None}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                validate_records([record(**changes)])
        validate_records([record(Credit=0)])

    def test_reconciliation_reports_mismatch_and_unavailable_separately(self):
        logger = logging.getLogger("test.reconcile")
        with patch("src.transform.validate.extract_summary_metrics", return_value={"total_debit": 100.49, "transaction_count": 1}):
            self.assertEqual(reconcile([record()], "sample.pdf", logger).status, "passed")
        with patch("src.transform.validate.extract_summary_metrics", return_value={"total_debit": 101, "transaction_count": 2}):
            result = reconcile([record()], "sample.pdf", logger)
            self.assertEqual(result.status, "failed")
            self.assertEqual(len(result.mismatches), 2)
        with patch("src.transform.validate.extract_summary_metrics", return_value={}):
            self.assertEqual(reconcile([record()], "sample.pdf", logger).status, "unavailable")

    def test_running_balance_check_supports_both_statement_orders(self):
        ascending = [record(Date='01/01/2026', Balance=900),
                     record(Date='02/01/2026', Debit=50, Balance=850)]
        self.assertEqual(check_running_balances(ascending).status, 'passed')
        ascending[1]['Balance'] = 800
        self.assertEqual(check_running_balances(ascending).status, 'mismatch')
        descending = list(reversed([record(Date='01/01/2026', Debit=50, Balance=950),
                                    record(Date='02/01/2026', Debit=50, Balance=900)]))
        self.assertEqual(check_running_balances(descending).status, 'passed')

    def test_cheque_notices_survive_validation_and_reconcile(self):
        logger = logging.getLogger("test.cheque_notice")
        for amount in (None, "", 0, 0.0):
            notice = record(Details="cheque rejected Chq:059010", Debit=amount, Credit=amount, Balance=0)
            rows = [record(Debit=100, Balance=900), notice, record(Debit=50, Balance=850)]
            validate_records(rows)
            self.assertEqual(check_running_balances(rows).status, "passed")
            for count in (2, 3):
                with patch("src.transform.validate.extract_summary_metrics", return_value={
                    "transaction_count": count, "total_debit": 150, "total_credit": 0,
                }):
                    self.assertEqual(reconcile(rows, "sample.pdf", logger).status, "passed")
            rows[-1]["Balance"] = 800
            result = check_running_balances(rows)
            self.assertEqual(result.status, "mismatch")
            self.assertIn("Rows 1-3", result.mismatches[0])
        with patch("src.transform.validate.extract_summary_metrics", return_value={"transaction_count": 4}):
            self.assertEqual(reconcile(rows, "sample.pdf", logger).status, "failed")

    def test_only_unmasked_identity_enables_deduplication(self):
        logger = logging.getLogger("test.identity")
        for account, expected in (("1234 567890", "axis:1234567890"), ("XXXX567890", None), ("", None)):
            with patch.object(main, "read_first_page", return_value=Mock(values={"Account Number": account})):
                self.assertEqual(main.account_identity(self.pdf, "axis", logger), expected)


class WorkbookSafetyTests(unittest.TestCase):
    def test_intermediate_export_sanitizes_text_preserves_precision_and_hides_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "output.xlsx"
            frame = records_to_dataframe([record(Details="=1+1\x06", _Account_Key="axis:123456789", _Source_Id="a.pdf")])
            write_output_excel(frame, target)
            wb = load_workbook(target, data_only=False)
            try:
                ws = wb["Statement"]
                headers = [cell.value for cell in ws[1]]
                self.assertFalse(any(h.startswith("_") for h in headers))
                self.assertEqual(ws["C2"].value, "=1+1")
                self.assertEqual(ws["C2"].data_type, "s")
                self.assertEqual(ws.cell(2, headers.index("Debit") + 1).value, 100.49)
            finally:
                wb.close()

    def test_repeated_same_second_names_never_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp, patch("src.utils.file_utils.datetime") as clock:
            clock.now.return_value = datetime(2026, 9, 17, 12, 34, 56)
            root = Path(tmp)
            expected = ["sample.xlsx", "sample_260917_123456.xlsx", "sample_260917_123456_1.xlsx", "sample_260917_123456_2.xlsx"]
            for name in expected:
                target = _next_final_path(root, "sample")
                self.assertEqual(target.name, name)
                target.write_bytes(name.encode())
            self.assertEqual((root / "sample.xlsx").read_bytes(), b"sample.xlsx")


class DetectionSafetyTests(unittest.TestCase):
    def test_generic_headings_do_not_identify_a_bank(self):
        for text in ("ACCOUNT STATEMENT", "ACCOUNT ACTIVITY", "STATEMENT OF ACCOUNT", "TRANSACTION ID", "ACCOUNT STATEMENT ACCOUNT ACTIVITY"):
            with self.subTest(text=text):
                self.assertIsNone(detector._detect_from_text(text))

    def test_tmb_is_identified_by_name_and_ifsc(self):
        for text in ("TAMILNAD MERCANTILE BANK ACCOUNT STATEMENT", "Tamilnadu Mercantile Bank", "IFSC TMBL0000123"):
            self.assertEqual(detector._detect_from_text(text), "tmb")

    def test_tied_bank_evidence_is_not_arbitrarily_selected(self):
        self.assertIsNone(detector._detect_from_text("AXIS BANK HDFC BANK"))

    def test_bank_names_nested_in_other_bank_names(self):
        for name, bank in (("SOUTH INDIAN BANK", "southind"), ("INDIAN BANK", "indian"), ("STATE BANK OF INDIA", "sbi"), ("CENTRAL BANK OF INDIA", "central"), ("CITY UNION BANK", "cub"), ("BANK OF INDIA", "boi")):
            with self.subTest(name=name):
                self.assertEqual(detector._detect_from_text(name), bank)

    def test_every_registered_bank_has_detection_signatures(self):
        self.assertEqual(set(PARSER_REGISTRY), set(detector.BANK_SIGNATURES))
        self.assertEqual(len(PARSER_REGISTRY), 24)


class DistributionTests(unittest.TestCase):
    def test_fresh_package_contains_installer_dependencies_and_runs_outside_checkout(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            repo = folder / "repo"
            repo.mkdir()
            for name in packager.ROOT_FILES_TO_INCLUDE:
                shutil.copy2(ROOT / name, repo / name)
            shutil.copytree(ROOT / "src", repo / "src", ignore=shutil.ignore_patterns("logs", "__pycache__", "*.bak"))
            (repo / "src" / "obsolete.py.bak").write_text("backup fixture", encoding="utf-8")
            (repo / "input").mkdir()
            (repo / "input" / "private.pdf").write_bytes(b"private fixture")
            archive = packager.create_package(repo, include_input_pdfs=False)
            with zipfile.ZipFile(archive) as bundle:
                names = bundle.namelist()
                self.assertTrue(any(name.endswith("/run.py") and name.count("/") == 1 for name in names))
                self.assertFalse(any(name.endswith("private.pdf") for name in names))
                self.assertFalse(any(name.endswith(".bak") for name in names))
                self.assertFalse(any("__pycache__" in name for name in names))
                bundle.extractall(folder / "unpacked")
            package = next((folder / "unpacked").iterdir())
            env = os.environ.copy()
            env["BANKSTMT_DB_PATH"] = str(folder / "history.db")
            env.pop("PYTHONPATH", None)
            for relative in ("run.py", "src/main.py", "src/code/run.py"):
                completed = subprocess.run([sys.executable, "-B", str(package / relative), "--help"], cwd=folder, env=env, capture_output=True, text=True, timeout=30)
                self.assertEqual(completed.returncode, 0, completed.stderr)
                self.assertIn("--pdf", completed.stdout)
            # A real native-text PDF exercises detection, extraction, validation,
            # precision and preservation of a legitimate repeated transaction.
            pdf_path = package / "input/cycle.pdf"
            with fitz.open() as doc:
                page = doc.new_page(width=650, height=500)
                for y, text in ((30, "Axis Bank"), (50, "Customer Name: TEST CUSTOMER"),
                                (70, "Account Number: 1234567890"), (90, "Address: Test Street"),
                                (110, "Statement Period: 01/01/2026 to 31/01/2026")):
                    page.insert_text((30, y), text, fontsize=10)
                xs = [30, 125, 350, 430, 510, 620]
                ys = [150, 180, 210, 240, 270]
                for x in xs:
                    page.draw_line((x, ys[0]), (x, ys[-1]))
                for y in ys:
                    page.draw_line((xs[0], y), (xs[-1], y))
                rows = [["Date", "Particulars", "Debit", "Credit", "Balance"],
                        ["01/01/2026", "ATM", "100.49", "", "899.51"],
                        ["01/01/2026", "CASH DEPOSIT", "", "100.49", "1000.00"],
                        ["01/01/2026", "ATM", "100.49", "", "899.51"]]
                for row_index, values in enumerate(rows):
                    for col, value in enumerate(values):
                        page.insert_text((xs[col] + 4, ys[row_index] + 18), value, fontsize=9)
                page.insert_text((30, 305), "Total debit 200.98 Total credit 100.49 Total transactions 3", fontsize=10)
                doc.save(pdf_path)
            pd.DataFrame(columns=["Category", "subCategory", "SheetName"]).to_excel(package / "input/Rules.xlsx", index=False)
            completed = subprocess.run([sys.executable, "-B", str(package / "run.py"), "--pdf", "cycle.pdf"], cwd=folder, env=env, capture_output=True, text=True, timeout=60)
            self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
            final_frame = pd.read_excel(package / "output/cycle.xlsx", sheet_name="Statement")
            self.assertEqual(len(final_frame), 3)
            self.assertEqual(final_frame["Debit"].tolist(), [100.49, 0., 100.49])
            self.assertEqual(final_frame["Balance"].tolist(), [899.51, 1000., 899.51])
            self.assertFalse(any(str(c).startswith("_") for c in final_frame.columns))
            # Load the legacy entry point, then build a real workbook without a
            # checkout on sys.path. This exercises the formerly failing imports.
            code = "import runpy, sys, logging; from pathlib import Path; runpy.run_path(sys.argv[1], run_name='legacy'); import pandas as pd; from src.export.final_excel_builder import build_final_workbook; build_final_workbook(pd.DataFrame(), Path('absent.xlsx'), Path(sys.argv[2]), 'smoke', logging.getLogger('smoke'))"
            completed = subprocess.run([sys.executable, "-B", "-c", code, str(package / "src/code/run.py"), str(folder / "results")], cwd=folder, env=env, capture_output=True, text=True, timeout=30)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertTrue((folder / "results/smoke.xlsx").is_file())


if __name__ == "__main__":
    unittest.main()
