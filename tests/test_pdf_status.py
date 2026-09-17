"""PDF_Status regressions using synthetic PDFs, without customer data."""
import hashlib
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import fitz
from openpyxl import Workbook

from src.export import final_excel_builder as builder
from src.transform.pdf_status import inspect_pdf, metadata_datetime
from src.utils.pdf_status_reader import read_first_page


def create_pdf(path, metadata=None, background=False, blank=False):
    with fitz.open() as doc:
        page = doc.new_page()
        if background:
            page.draw_rect(fitz.Rect(30, 30, 450, 180), fill=(1, 1, 1), overlay=False)
        if not blank:
            for y, text in ((50, 'Axis Bank'), (75, 'Customer Name: TEST TRADERS'),
                            (95, 'Account Number: 000123456789'), (115, 'Address: 12 TEST ROAD'),
                            (135, 'IFSC: UTIB0001234'), (155, 'Statement period: 01/01/2026 to 31/01/2026'),
                            (190, 'Date Particulars Debit Credit Balance'),
                            (215, '01/01/2026 TEST 100.00 0.00 900.00')):
                page.insert_text((40, y), text, fontsize=10)
        if metadata:
            doc.set_metadata(metadata)
        doc.save(path)


def audit(path, password=None):
    return inspect_pdf(path, password, first_page=read_first_page(path, password, allow_ocr=False))


class PDFStatusTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.path = Path(self.directory.name) / 'statement.pdf'

    def tearDown(self):
        self.directory.cleanup()

    def test_unmodified_is_no_evidence_not_proven_authentic(self):
        create_pdf(self.path)
        before = hashlib.sha256(self.path.read_bytes()).digest()
        rows = audit(self.path)
        self.assertEqual(rows[0]['Result'], 'No evidence of manual editing found')
        self.assertEqual(rows[0]['Status'], 'PASS')
        self.assertIn('cannot', rows[0]['Details'])
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).digest(), before)

    def test_all_five_fields_from_first_page(self):
        create_pdf(self.path)
        self.assertEqual(read_first_page(self.path, allow_ocr=False).values, {
            'Customer Name': 'TEST TRADERS', 'Bank Name': 'Axis Bank',
            'Account Number': '000123456789', 'Address': '12 TEST ROAD',
            'Statement Date Between': '2026-01-01 to 2026-01-31'})

    def test_content_change_between_revisions(self):
        create_pdf(self.path)
        with fitz.open(self.path) as doc:
            doc[0].insert_text((40, 240), 'Extra transaction 5000.00')
            doc.saveIncr()
        rows = audit(self.path)
        self.assertEqual(rows[0]['Status'], 'FAIL')
        self.assertEqual(rows[0]['Result'], 'Content changes detected after an earlier save')
        detail = next(r['Details'] for r in rows if r['Check'] == 'Saved revision content comparison')
        self.assertIn('Page 1', detail)
        self.assertIn('5000.00', detail)

    def test_metadata_only_save_is_not_content_edit(self):
        create_pdf(self.path)
        with fitz.open(self.path) as doc:
            doc.set_metadata({'title': 'Statement copy'})
            doc.saveIncr()
        self.assertEqual(audit(self.path)[0]['Status'], 'PASS')

    def test_reverted_intermediate_edit_is_detected(self):
        create_pdf(self.path)
        with fitz.open(self.path) as doc:
            original = doc[0].get_contents()
            doc[0].insert_text((40, 240), 'Temporary edit')
            doc.saveIncr()
        with fitz.open(self.path) as doc:
            doc.xref_set_key(doc[0].xref, 'Contents', '[' + ' '.join(f'{x} 0 R' for x in original) + ']')
            doc.saveIncr()
        self.assertEqual(audit(self.path)[0]['Status'], 'FAIL')

    def test_marker_strings_are_not_revisions_or_javascript(self):
        create_pdf(self.path)
        with fitz.open(self.path) as doc:
            doc[0].insert_text((40, 245), '%%EOF /Prev /JavaScript /JS /Sig startxref')
            replacement = self.path.with_name('markers.pdf')
            doc.save(replacement)
        self.assertEqual(audit(replacement)[0]['Status'], 'PASS')

    def test_background_is_not_coverup(self):
        create_pdf(self.path, background=True)
        rows = audit(self.path)
        self.assertEqual(next(r['Status'] for r in rows if r['Check'] == 'Text overlays and cover-ups'), 'PASS')

    def test_flattened_coverup_without_revision_history(self):
        create_pdf(self.path)
        with fitz.open(self.path) as doc:
            doc[0].draw_rect(fitz.Rect(38, 202, 330, 220), fill=(1, 1, 1), color=(1, 1, 1))
            doc[0].insert_text((40, 215), '01/01/2026 TEST 900.00 0.00 100.00', fontsize=10)
            replacement = self.path.with_name('overlay.pdf')
            doc.save(replacement, garbage=4)
        rows = audit(replacement)
        self.assertEqual(rows[0]['Result'], 'Possible manual editing - review evidence')
        self.assertEqual(next(r['Status'] for r in rows if r['Check'] == 'Text overlays and cover-ups'), 'WARNING')

    def test_later_pages_inspected(self):
        create_pdf(self.path)
        with fitz.open(self.path) as doc:
            doc.new_page().insert_text((40, 50), 'Hidden data', render_mode=3)
            doc.saveIncr()
        self.assertIn('Page 2', next(r['Details'] for r in audit(self.path) if r['Check'] == 'Hidden text'))

    def test_unknown_password_is_not_modification(self):
        create_pdf(self.path)
        encrypted = self.path.with_name('encrypted.pdf')
        with fitz.open(self.path) as doc:
            doc.save(encrypted, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='secret')
        self.assertIn('Cannot assess', audit(encrypted)[0]['Result'])
        self.assertEqual(audit(encrypted, 'secret')[0]['Result'], 'No evidence of manual editing found')

    def test_missing_corrupt_files_not_reported_as_edited(self):
        self.assertIn('Cannot assess', audit(self.path)[0]['Result'])
        self.path.write_bytes(b'not a PDF')
        self.assertIn('Cannot assess', audit(self.path)[0]['Result'])

    def test_missing_required_fields_are_explicit(self):
        create_pdf(self.path, blank=True)
        rows = audit(self.path)
        self.assertIn('Inconclusive', rows[0]['Result'])
        self.assertIn('Account Number', next(r['Result'] for r in rows if r['Check'] == 'Required first-page details'))

    def test_second_page_and_filename_cannot_supply_identity(self):
        self.path = self.path.with_name('Axis Bank 123456789.pdf')
        create_pdf(self.path, blank=True)
        with fitz.open(self.path) as doc:
            doc.new_page().insert_text((40, 50), 'Axis Bank Customer Name: WRONG PERSON Account No: 123456789')
            doc.saveIncr()
        self.assertFalse(any(read_first_page(self.path, allow_ocr=False).values.values()))

    def test_equivalent_metadata_timezones(self):
        self.assertEqual(metadata_datetime("D:20260101153000+05'30'"), metadata_datetime('D:20260101100000Z'))
        self.assertEqual(metadata_datetime('2026-01-01T15:30:00+05:30'), metadata_datetime('D:20260101100000Z'))
        create_pdf(self.path, {'creationDate': "D:20260101153000+05'30'", 'modDate': 'D:20260101100000Z'})
        self.assertEqual(next(r['Status'] for r in audit(self.path) if r['Check'] == 'Metadata dates'), 'PASS')

    def test_unknown_metadata_timezone_requires_review(self):
        create_pdf(self.path, {'creationDate': 'D:20260101100000', 'modDate': 'D:20260101100000Z'})
        self.assertEqual(next(r['Status'] for r in audit(self.path) if r['Check'] == 'Metadata dates'), 'WARNING')

    def test_editor_metadata_not_confirmed_manual_edit(self):
        create_pdf(self.path, {'creator': 'Microsoft Word'})
        self.assertEqual(audit(self.path)[0]['Result'], 'Possible manual editing - review evidence')

    def test_ocr_unavailable_reported(self):
        create_pdf(self.path, blank=True)
        with patch('src.utils.pdf_status_reader.ocr_words', side_effect=RuntimeError('OCR unavailable')):
            result = read_first_page(self.path)
        self.assertIn('OCR unavailable', result.notes[0])
        self.assertFalse(result.values['Bank Name'])

    def test_two_columns_and_masked_account(self):
        with fitz.open() as doc:
            page = doc.new_page()
            for x, y, text in ((40, 40, 'Bank of Baroda'), (40, 70, 'Branch Address: BANK ROAD'),
                               (40, 105, 'A/C Name: ALPHA TRADERS'), (340, 105, 'Account Type: Current'),
                               (40, 125, 'Address: 21 CUSTOMER ROAD'), (340, 125, 'IFSC Code: BARB0ABCDEF'),
                               (40, 145, 'A/C Number: 000XXXX789'),
                               (40, 170, 'Statement for period 01 Jan 2026 - 31 Jan 2026'),
                               (40, 200, 'Date Description Debit Credit Balance'),
                               (40, 225, '01/01/2026 Indian Bank 100.00 0.00 900.00')):
                page.insert_text((x, y), text, fontsize=10)
            doc.save(self.path)
        values = read_first_page(self.path, allow_ocr=False).values
        self.assertEqual(values['Bank Name'], 'Bank of Baroda')
        self.assertEqual(values['Customer Name'], 'ALPHA TRADERS')
        self.assertEqual(values['Account Number'], '000XXXX789')
        self.assertEqual(values['Address'], '21 CUSTOMER ROAD')

    def image_bytes(self):
        pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 10, 10), False)
        pixmap.clear_with(255)
        return pixmap.tobytes('png')

    def save_rewrite(self, modify):
        create_pdf(self.path)
        rewritten = self.path.with_name('rewritten.pdf')
        with fitz.open(self.path) as doc:
            modify(doc)
            doc.save(rewritten, garbage=4, deflate=True)
        return rewritten

    def test_rewritten_image_coverup_is_detected(self):
        path = self.save_rewrite(lambda doc: doc[0].insert_image(
            fitz.Rect(38, 202, 330, 220), stream=self.image_bytes(), keep_proportion=False))
        rows = audit(path)
        self.assertEqual(rows[0]['Result'], 'Possible manual editing - review evidence')
        evidence = next(r for r in rows if r['Check'] == 'Image overlays')
        self.assertEqual(evidence['Status'], 'WARNING')
        self.assertIn('Page 1', evidence['Details'])

    def test_image_background_preceding_text_is_not_coverup(self):
        path = self.save_rewrite(lambda doc: doc[0].insert_image(
            fitz.Rect(38, 202, 330, 220), stream=self.image_bytes(), keep_proportion=False, overlay=False))
        self.assertEqual(audit(path)[0]['Status'], 'PASS')

    def test_conflicting_amounts_without_rectangle(self):
        def modify(doc):
            page = doc[0]
            rect = page.search_for('100.00')[0]
            page.insert_text((rect.x0, 215), '900.00', fontsize=10)
        rows = audit(self.save_rewrite(modify))
        evidence = next(r for r in rows if r['Check'] == 'Conflicting amount overprints')
        self.assertEqual(evidence['Status'], 'WARNING')
        self.assertIn("'100.00' and '900.00'", evidence['Details'])

    def test_duplicate_amount_paint_is_not_edit_evidence(self):
        def modify(doc):
            page = doc[0]
            rect = page.search_for('100.00')[0]
            page.insert_text((rect.x0, 215), '100.00', fontsize=10)
        self.assertEqual(audit(self.save_rewrite(modify))[0]['Status'], 'PASS')

    def test_large_image_with_readable_header_is_inconclusive(self):
        path = self.save_rewrite(lambda doc: doc[0].insert_image(
            fitz.Rect(0, 250, 595, 842), stream=self.image_bytes(), keep_proportion=False))
        rows = audit(path)
        self.assertEqual(rows[0]['Result'], 'Inconclusive - inspection has limitations')
        self.assertIn('Large raster images', rows[0]['Details'])

    def test_xmp_only_editor_is_reported(self):
        path = self.save_rewrite(lambda doc: doc.set_xml_metadata(
            '<x:xmpmeta xmlns:x="adobe:ns:meta/" xmlns:xmp="http://ns.adobe.com/xap/1.0/">'
            '<xmp:CreatorTool>Microsoft Word</xmp:CreatorTool></x:xmpmeta>'))
        self.assertIn('Editing software recorded in XMP', audit(path)[0]['Details'])

    def test_invalid_metadata_date_requires_review(self):
        create_pdf(self.path, {'creationDate': 'D:20261399999999', 'modDate': 'D:20260101000000Z'})
        rows = audit(self.path)
        self.assertEqual(next(r['Result'] for r in rows if r['Check'] == 'Metadata dates'), 'Invalid metadata date')
        self.assertEqual(rows[0]['Status'], 'WARNING')

    def test_failure_does_not_erase_confirmed_revision_change(self):
        create_pdf(self.path)
        with fitz.open(self.path) as doc:
            doc[0].insert_text((40, 240), 'Extra transaction 5000.00')
            doc.saveIncr()
        for target in ('page_evidence', 'object_evidence'):
            with self.subTest(target=target), patch('src.transform.pdf_status.' + target, side_effect=RuntimeError('failed')):
                rows = audit(self.path)
            self.assertEqual(rows[0]['Status'], 'FAIL')
            self.assertIn('Limitations:', rows[0]['Details'])
            self.assertTrue(any(r['Check'] == 'Saved revision content comparison' for r in rows))

    def test_failed_revision_render_preserves_text_change(self):
        create_pdf(self.path)
        with fitz.open(self.path) as doc:
            doc[0].insert_text((40, 240), 'Extra transaction 5000.00')
            doc.saveIncr()
        with patch('src.transform.pdf_status._page_digest', side_effect=RuntimeError('render failed')):
            rows = audit(self.path)
        self.assertEqual(rows[0]['Status'], 'FAIL')
        self.assertIn('comparison failed', rows[0]['Details'])

    def test_failed_page_check_cannot_pass(self):
        create_pdf(self.path)
        with patch('src.transform.pdf_status.page_evidence', side_effect=RuntimeError('failed')):
            rows = audit(self.path)
        self.assertEqual(rows[0]['Status'], 'WARNING')
        self.assertEqual(next(r['Status'] for r in rows if r['Check'] == 'Image overlays'), 'WARNING')

    def test_revision_limit_is_explicit(self):
        create_pdf(self.path)
        for title in ('copy1', 'copy2'):
            with fitz.open(self.path) as doc:
                doc.set_metadata({'title': title})
                doc.saveIncr()
        with patch('src.transform.pdf_status.MAX_REVISION_CANDIDATES', 1):
            rows = audit(self.path)
        self.assertEqual(rows[0]['Status'], 'WARNING')
        self.assertIn('limit reached', rows[0]['Details'])

    def test_first_page_failure_does_not_skip_forensics(self):
        create_pdf(self.path)
        with patch('src.transform.pdf_status.read_first_page', side_effect=RuntimeError('failed')):
            rows = inspect_pdf(self.path)
        self.assertEqual(rows[0]['Status'], 'WARNING')
        self.assertTrue(any(r['Check'] == 'Text overlays and cover-ups' for r in rows))

    def signature_fixture(self, byte_range=None):
        create_pdf(self.path)
        path = self.path.with_name('signature.pdf')
        with fitz.open(self.path) as doc:
            xref = doc.get_new_xref()
            doc.update_object(xref, '<< /Type /Sig /Contents <30820000> /ByteRange '
                              '[0 1111111111 2222222222 3333333333] >>')
            doc.xref_set_key(doc.pdf_catalog(), 'TestSignature', f'{xref} 0 R')
            doc.save(path)
        raw = path.read_bytes()
        start = raw.index(b'<30820000>')
        end = start + len(b'<30820000>')
        values = byte_range or f'[0 {start:010d} {end:010d} {len(raw) - end:010d}]'
        old = b'[0 1111111111 2222222222 3333333333]'
        self.assertEqual(len(values), len(old))
        self.assertIn(old, raw)
        path.write_bytes(raw.replace(old, values.encode('ascii')))
        return path

    def test_unverified_signature_does_not_make_overall_pass(self):
        rows = audit(self.signature_fixture())
        self.assertEqual(rows[0]['Status'], 'WARNING')
        self.assertIn('cryptographic verification', rows[0]['Details'])
        evidence = next(r for r in rows if r['Check'] == 'Digital signature')
        self.assertIn('structural range valid', evidence['Details'])

    def test_signature_decimal_range_is_invalid(self):
        rows = audit(self.signature_fixture('[0 111111.111 2222222222 3333333333]'))
        evidence = next(r for r in rows if r['Check'] == 'Digital signature')
        self.assertIn('structural range invalid', evidence['Details'])

    def test_signed_revision_trailing_bytes_are_reported(self):
        path = self.signature_fixture()
        with fitz.open(path) as doc:
            doc.set_metadata({'title': 'later save'})
            doc.saveIncr()
        rows = audit(path)
        self.assertIn('bytes after a signed revision', rows[0]['Details'])

    def test_image_coverup_on_later_page_after_full_rewrite(self):
        def modify(doc):
            page = doc.new_page()
            page.insert_text((40, 215), 'Payment 100.00', fontsize=10)
            page.insert_image(fitz.Rect(38, 202, 330, 220), stream=self.image_bytes(), keep_proportion=False)
        rows = audit(self.save_rewrite(modify))
        evidence = next(r for r in rows if r['Check'] == 'Image overlays')
        self.assertEqual(evidence['Status'], 'WARNING')
        self.assertIn('Page 2', evidence['Details'])

    def test_image_only_revision_change_is_confirmed(self):
        create_pdf(self.path)
        with fitz.open(self.path) as doc:
            doc[0].draw_rect(fitz.Rect(400, 400, 450, 450), fill=(0, 0, 0))
            doc.saveIncr()
        rows = audit(self.path)
        self.assertEqual(rows[0]['Status'], 'FAIL')
        self.assertIn('appearance changed', next(r['Details'] for r in rows if r['Check'] == 'Saved revision content comparison'))

    def test_zero_signature_placeholder_is_not_valid(self):
        path = self.signature_fixture()
        path.write_bytes(path.read_bytes().replace(b'<30820000>', b'<00000000>'))
        rows = audit(path)
        self.assertIn('structural range invalid', next(r['Details'] for r in rows if r['Check'] == 'Digital signature'))

    def test_signature_without_byte_range_is_reported(self):
        def modify(doc):
            xref = doc.get_new_xref()
            doc.update_object(xref, '<< /Type /Sig /Contents <30820000> >>')
            doc.xref_set_key(doc.pdf_catalog(), 'TestSignature', f'{xref} 0 R')
        rows = audit(self.save_rewrite(modify))
        self.assertEqual(rows[0]['Status'], 'WARNING')
        self.assertIn('structural range invalid', next(r['Details'] for r in rows if r['Check'] == 'Digital signature'))

    def test_new_evidence_reaches_final_sheet(self):
        path = self.save_rewrite(lambda doc: doc[0].insert_image(
            fitz.Rect(38, 202, 330, 220), stream=self.image_bytes(), keep_proportion=False))
        frame = builder._build_pdf_status_sheet([path])
        self.assertEqual(list(frame.columns), builder.PDF_STATUS_COLUMNS)
        self.assertEqual(frame.loc[frame['Check'] == 'Image overlays', 'Status'].iloc[0], 'WARNING')

    def test_style_does_not_modify_other_sheet(self):
        book = Workbook()
        other = book.active
        other.title = 'Statement'
        other['A1'] = '=SUM(B1:B2)'
        other['B1'] = 123.45
        other.column_dimensions['A'].width = 29
        status = book.create_sheet('PDF_Status', 0)
        for col, label in enumerate(builder.PDF_STATUS_COLUMNS, 1):
            status.cell(8, col, label)
        status.cell(9, 3, 'WARNING')
        status.cell(9, 5, 'Evidence\n' * 8)
        builder._apply_pdf_status_style(book, 'PDF_Status', [(label, 'sample') for label in builder.PDF_ACCOUNT_SUMMARY_LABELS])
        self.assertEqual(other['A1'].value, '=SUM(B1:B2)')
        self.assertEqual(other['B1'].value, 123.45)
        self.assertEqual(other.column_dimensions['A'].width, 29)
        self.assertGreater(status.row_dimensions[9].height, 36)
        self.assertEqual(status.freeze_panes, 'A9')


if __name__ == '__main__':
    unittest.main()
