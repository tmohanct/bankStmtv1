from __future__ import annotations

import logging
import re
import shutil
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src" / "code"))

import final_excel_builder
from utils import OUTPUT_COLUMNS, clean_detail


ILLEGAL_CONTROL_CHAR_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


class ExcelIllegalCharacterTests(unittest.TestCase):
    def test_final_workbook_removes_illegal_control_characters(self) -> None:
        details = "Transfer\x00 with\x06 invalid\x1f metadata"
        statement_df = pd.DataFrame(
            [
                {
                    "Sno": 1,
                    "Date": "01/01/2026",
                    "Details": details,
                    "Detail_Clean": clean_detail(details),
                    "Cheque No": "",
                    "Debit": 100.0,
                    "Credit": 0.0,
                    "Balance": 900.0,
                    "Source": "sample.pdf",
                }
            ],
            columns=OUTPUT_COLUMNS,
        )
        pdf_status_df = pd.DataFrame(
            [
                {
                    "PDF": "sample.pdf",
                    "Check": "PDF creator",
                    "Status": "PASS",
                    "Result": "Creator metadata found",
                    "Details": "Creator:\nProducer: ð\x06¬Ï",
                }
            ],
            columns=final_excel_builder.PDF_STATUS_COLUMNS,
        )
        account_summary_rows = [
            ("Customer Name", "ARUN\x06 HOSPITAL"),
            ("Bank Name", "HDFC Bank"),
            ("Account Number", "1234"),
            ("Address", "Test address"),
            ("Statement Date Between", "01/01/2026 to 31/01/2026"),
        ]

        logger = logging.getLogger("tests.excel_illegal_characters")
        logger.handlers.clear()
        logger.addHandler(logging.NullHandler())

        temp_root = PROJECT_ROOT / "output" / "_excel_illegal_characters_test"
        shutil.rmtree(temp_root, ignore_errors=True)
        temp_root.mkdir(parents=True, exist_ok=True)

        try:
            rules_path = temp_root / "Rules.xlsx"
            output_dir = temp_root / "output"
            pd.DataFrame(columns=["Category", "subCategory", "SheetName"]).to_excel(
                rules_path,
                index=False,
            )

            with (
                patch.object(final_excel_builder, "_build_pdf_status_sheet", return_value=pdf_status_df),
                patch.object(
                    final_excel_builder,
                    "_build_pdf_account_summary_rows",
                    return_value=account_summary_rows,
                ),
            ):
                final_path = final_excel_builder.build_final_workbook(
                    statement_df=statement_df,
                    rules_path=rules_path,
                    output_dir=output_dir,
                    pdf_stem="sample",
                    logger=logger,
                )

            workbook = load_workbook(final_path, data_only=False)
            try:
                for worksheet in workbook.worksheets:
                    for row in worksheet.iter_rows():
                        for cell in row:
                            if isinstance(cell.value, str):
                                self.assertIsNone(
                                    ILLEGAL_CONTROL_CHAR_RE.search(cell.value),
                                    f"Illegal character remained in {worksheet.title}!{cell.coordinate}",
                                )

                self.assertEqual(workbook["PDF_Status"]["B1"].value, "ARUN HOSPITAL")
                self.assertEqual(
                    workbook["PDF_Status"]["E9"].value,
                    "Creator:\nProducer: ð¬Ï",
                )
                self.assertEqual(
                    workbook["Statement"]["C2"].value,
                    "Transfer with invalid metadata",
                )
            finally:
                workbook.close()
        finally:
            shutil.rmtree(temp_root, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
